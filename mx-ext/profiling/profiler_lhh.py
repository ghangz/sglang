#start the code by lhh
import os
import pandas as pd
import glob
import numpy as np
import copy
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, numbers

def filter_data(data_in):
    c500_data = pd.read_excel(data_in)
    name_list = c500_data['Instances']
    total_list = c500_data['Total Time (us)']
    avg_list = c500_data['Avg (us)']
 
    time_list = c500_data['Time(%)']
    prefill_data = c500_data[c500_data['stage']=='prefill']
    decoding_data = c500_data[c500_data['stage']=='decoding'] 

    key_word = ['triton','_fwd_kernel_stage','sglang','vllm','mccl','nccl','cutlass','gemm','Tokens','Requests', 'moe_align_block_size_kernel','elementwise_kernel','mbtopk::computeBlockwiseKthCounts',"RMSNormKernel" ]
    ind = '|'.join(key_word)

    prefill_a  = prefill_data[ prefill_data['Name'].str.contains(ind) | prefill_data['Time(%)'].map(lambda x : float(x.strip('%')) >1)]
    decoding_a = decoding_data[ decoding_data['Name'].str.contains(ind) | decoding_data['Time(%)'].map(lambda x : float(x.strip('%')) >1)]
    prefill_differ =list(set(prefill_data['Name']) - set(prefill_a['Name']))
    print(f'Prefill differ: {prefill_differ}')
    decode_differ =list(set(decoding_data['Name']) - set(decoding_a['Name']))
    print(' ')
    print(f'Decode differ: {decode_differ}')
    return prefill_a, decoding_a


def merge_key(data_prefill, keys):
    new_df = {}
    ins_cont = 0
    total_cont = 0
    time_cont = 0

    name_list = []
    ins_list = []
    total_list = []
    time_list = []

    for name in data_prefill['Name']:
        idx = data_prefill['Name'].tolist().index(name)
        ins = data_prefill['Instances'].to_list()[idx]
        total_time = data_prefill['Total Time (us)'].to_list()[idx]
        time = data_prefill['Time(%)'].to_list()[idx]

        if keys in name:
            ins_cont = ins_cont + ins
            total_cont = total_cont + total_time
            time_cont = time_cont + float(str(time).strip('%'))
        else:
            name_list.append(name)
            ins_list.append(ins)
            total_list.append(total_time)
            time_list.append(str(time).strip('%'))
    name_list.insert(0,keys)
    ins_list.insert(0,ins_cont)
    total_list.insert(0,total_cont)
    time_list.insert(0,time_cont)
    
    new_df['Name'] = name_list
    new_df['Instances'] = ins_list
    new_df['Total Time (us)'] = total_list
    new_df['Time(%)'] = time_list

    return new_df


def mergea100_c500(c500_pre, a100_pre):
    final_set = list(set(c500_pre['Name']) | set(a100_pre['Name']))

    total_name_list = []
    c500_total_ins_list = []
    c500_total_time_list=[]
    c500_total_per_list = []


    a100_total_ins_list = []
    a100_total_time_list=[]
    a100_total_per_list = []

    for fi in final_set:
        total_name_list.append(fi)
        try:
            c500_idx = c500_pre['Name'].index(fi)

            c500_total_ins_list.append(c500_pre['Instances'][c500_idx])
            c500_total_time_list.append(c500_pre['Total Time (us)'][c500_idx])
            c500_total_per_list.append(c500_pre['Time(%)'][c500_idx])
        except:
            c500_total_ins_list.append('Nan')
            c500_total_time_list.append('Nan')
            c500_total_per_list.append('Nan')

        try:
            a100_idx = a100_pre['Name'].index(fi)

            # a100_total_name_list.append(fi)/
            a100_total_ins_list.append(a100_pre['Instances'][a100_idx])
            a100_total_time_list.append(a100_pre['Total Time (us)'][a100_idx])
            a100_total_per_list.append(a100_pre['Time(%)'][a100_idx])
        except:
            # a100_total_name_list.append(fi)
            a100_total_ins_list.append('Nan')
            a100_total_time_list.append('Nan')
            a100_total_per_list.append('Nan')
    return total_name_list,c500_total_ins_list,c500_total_time_list,c500_total_per_list,a100_total_ins_list,a100_total_time_list,a100_total_per_list


def data_nan_to_zero(s):
    if isinstance(s, str) and s.lower()=="nan":
        return 0
    try:
        int_val = int(s)
        return int_val
    except ValueError:
        try:
            float_val = float(s)
            return float_val
        except ValueError:
            return 0

def get_summary(data_process):
    col_info =  {"time_c500":0, "call_c500":0 ,"percent_c500":0, 
                "time_a100":0, "call_a100":0 ,"percent_a100":0, "pecent(a100/c500)":0}
    #LIUHH
    name_row = ["gemm", 
                "fused-moe", 
                "mccl",
                "moe_align_block_size",
                "elementwise_kernel",
                "at::native::mbtopk",
                "grouped_gemm_triton_kerne",
                "layernorm", 
                "act_and_mul",
                "moe_align_block_size_kernel",
                "at::native::sbtopk::gatherTopK",
                "_fwd_kerne",
                "_fwd_kernel_stage",
                "_fwd_grouped_kernel_stage", 
                "at::native::bitonicSortKVInPlace",
                "RMSNormKernel",
                "vllm::dynamic_scaled_int8_quant_kernel",
                "triton",
                "other",
                "Tokens",
                "Requests"]
    summary_init = {
    }

    for name in name_row:
        summary_init[name] = copy.deepcopy(col_info)

    summary = copy.deepcopy(summary_init)
    
    name_mapping = {
        # c500 mapping
        # new
        "write_req_to_token_pool_trito":"triton",
        "at_cuda_detail::cub::DeviceScanInitKernel":"at_cuda_detail::cub::DeviceScanInitKernel",
        "_fwd_kerne":"_fwd_kerne",
        "at::native::sbtopk::gatherTopK":"at::native::sbtopk::gatherTopK",
        "at_cuda_detail::cub::DeviceScanKernel":"at_cuda_detail::cub::DeviceScanKernel",
        "gemm" : "gemm",
        "cutlass::Kernel" : "gemm",     #LIUHH
        "cublasLt::splitKreduce_kernel" : "gemm",
        # "ampere_fp16_s16816gemm" : "gemm",
        "mcblas::DevDstElementwiseAdd" : "gemm",
        "dot_kernel": "gemm",
        "reduce_1Block_kernel": "gemm",
        # "gemv2T_kernel_val": "gemm",
        # "internal::gemvx::kernel": "gemm",
        # "gemvnn_splitk" : "gemm",
        "mcblas__Mck_hge" : "gemm",

        "paged_attention" : "vllm:paged_attn",
        "rotary_embedding_kernel" : "vllm:rotary_embeded",
        "reshape_and_cache_kernel" : "vllm:reshape_and_cache",

        "triton":"triton",
        "grouped_gemm_triton_kerne":"grouped_gemm_triton_kerne",
        "act_and_mul_kernel" : "act_and_mul",
        "cunn_SoftMaxForward" : "Softmax",

        "fused_moe_kerne" : "fused-moe",
        "fusedMoe" : "fused-moe",
        "mcblas__Mck_fp16_fusedMoe_": "fused-moe",

        "Tokens"  : "Tokens",
        "Requests" : "Requests",
        "mla" : "MLA",
        "_fwd_grouped_kernel_stage" : "_fwd_grouped_kernel_stage",
        "_fwd_kernel_stage" : "_fwd_kernel_stage",

        "fused_add_rms_norm_kernel" : "layernorm",
        "enable_if" : "layernorm",
        "at::native::reduce_kernel":"layernorm",
        "rms_norm_kernel" : "layernorm",     

        "mccl" : "mccl",
        "mcclKernel" : "mccl",
        "ncclKernel" : "mccl",
        "nccl" : "mccl",
        "cross_device_reduce_1stage" : "mccl",
        "cross_device_reduce_2stage" : "mccl",
        "cross_device_reduce_3stage" : "mccl",
        "RMSNormKernel":"RMSNormKernel",
        "Tokens"  : "Tokens",
        "Requests" : "Requests",
        "reshape_and_cache_flash_kernel" : "vllm:reshape_and_cache", 
        "moe_align_block_size_kernel" : "moe_align_block_size",
        "compute_position_kerne":"compute_position_kerne",
        
        # "at::native::vectorized_elementwise_kernel":"elementwise_kernel",
        # "(anonymous namespace)::elementwise_kernel_with_index":"elementwise_kernel",
        # "at::native::unrolled_elementwise_kernel":"elementwise_kernel",
        # "at::native::_scatter_gather_elementwise_kernel":"elementwise_kernel",
        "elementwise_kernel":"elementwise_kernel", 
        "CatArrayBatchedCopy":"elementwise_kernel",
        
        "at::native::sbtopk::gatherTopK":"at::native::sbtopk::gatherTopK",
        "at::native::bitonicSortKVInPlace" : "at::native::bitonicSortKVInPlace", #LIUHH-C500       
        
        "at::native::mbtopk::computeBlockwiseKthCounts" : "at::native::mbtopk", #LIUHH
        "at::native::mbtopk::gatherTopK" : "at::native::mbtopk",#LIUHH
        "at::native::mbtopk::computeBlockwiseKthCounts" : "at::native::mbtopk",#LIUHH
        "at::native::mbtopk::fill" : "at::native::mbtopk",#LIUHH
        "at::native::mbtopk::radixFindKthValues" : "at::native::mbtopk",#LIUHH
        # "mcblas::DevDstElementwiseAdd" : "mcblas::DevDstElementwiseAdd",#LIUHH
        "at::native::InputPerOutputImcontinuousReduceKernel" : "at::native::InputPerOutputImcontinuousReduceKernel",#LIUHH
        "vllm::dynamic_scaled_int8_quant_kernel" : "vllm::dynamic_scaled_int8_quant_kernel",
    }


    for _, row in data_process.iterrows():
        name = row['Name']
        call_of_c500 = data_nan_to_zero(row["c500_instance"])
        time_of_c500 = data_nan_to_zero(row["c500_time"])
        percent_of_c500 = data_nan_to_zero(row["c500_per"])

        call_of_a100 = data_nan_to_zero(row["a100_instance"])
        time_of_a100 = data_nan_to_zero(row["a100_time"])
        percent_of_a100 = data_nan_to_zero(row["a100_per"])

        if "grouped_gemm_triton_kerne" in name or "triton" == name:
            kernel_name = name
            summary[kernel_name]["time_c500"] += time_of_c500
            summary[kernel_name]["call_c500"] += call_of_c500
            summary[kernel_name]["percent_c500"] += percent_of_c500

            summary[kernel_name]["time_a100"] += time_of_a100
            summary[kernel_name]["call_a100"] += call_of_a100
            summary[kernel_name]["percent_a100"] += percent_of_a100
            continue

        flag = False
        for prefix_name, kernel_name in name_mapping.items():
            if name.find(prefix_name) >= 0:
                # import pdb
                #pdb.set_trace()
                # print(kernel_name)
                # print(prefix_name)
                if(name=='Requests'or name=='Tokens'):
                    print(time_of_c500)
                    print(call_of_c500)
                    print(row)
                    #pdb.set_trace()
                summary[kernel_name]["time_c500"] += time_of_c500
                if prefix_name != "mcblas::DevDstElementwiseAdd":
                    summary[kernel_name]["call_c500"] += call_of_c500
                else:
                    print("Shot mcblas::DevDstElementwiseAdd")
                summary[kernel_name]["percent_c500"] += percent_of_c500

                summary[kernel_name]["time_a100"] += time_of_a100
                summary[kernel_name]["call_a100"] += call_of_a100
                summary[kernel_name]["percent_a100"] += percent_of_a100
                # print(summary[kernel_name]["percent_a100"])

                flag = True
        
        if not flag:
            print(f"Other kernel: {name}")
            kernel_name = "other"
            summary[kernel_name]["time_c500"] += time_of_c500
            summary[kernel_name]["call_c500"] += call_of_c500
            summary[kernel_name]["percent_c500"] += percent_of_c500

            summary[kernel_name]["time_a100"] += time_of_a100
            summary[kernel_name]["call_a100"] += call_of_a100
            summary[kernel_name]["percent_a100"] += percent_of_a100
    
    for key in summary.keys():
        if summary[key]["time_c500"] < 1e-5:
            summary[key]["pecent(a100/c500)"]  = "NaN"
        else:
            summary[key]["pecent(a100/c500)"] =  summary[key]["time_a100"]/summary[key]["time_c500"]
    
    return summary



def process(data, model_name="default"):
    data_prefill = data[data['stage'] == 'prefill']
    data_decoding = data[data['stage'] == 'decoding']


    summary_prefill = get_summary(data_prefill)
    summary_decoding = get_summary(data_decoding)
    all_data_dict = {}
    all_data_dict["prefill"] = summary_prefill
    all_data_dict["decoding"] = summary_decoding


    rows = []
    first_model = True
    for stage, kernels in all_data_dict.items():
        first_kernel = True
        for kernel_name, metrics in kernels.items():
            
            if first_kernel:
                row = {'model': model_name,'stage': stage, 'kernel_name': kernel_name}
                first_kernel = True
            else:
                row = {'model': model_name, 'stage': "", 'kernel_name': kernel_name}
            if first_model:
                model_name = ""
                first_kernel = True
            row.update(metrics)
            rows.append(row)
    return rows



def get_concurrency(dataset_path,a100_path):

    all_rows = []
    print(dataset_path)
    files = [f for f in os.listdir(dataset_path) if f.endswith('.xlsx')]
    files.sort()
    files_a100 = [f for f in os.listdir(a100_path) if f.endswith('.xlsx')]
    files_a100.sort()
    print("assert")
    print(len(files))
    print(len(files_a100))
    assert (len(files))==(len(files_a100)),"files number is not equal"
    print(len(files))
    print(len(files_a100))
    # one_file_data = pd.read_excel("/AI-DATA/job-lhh/all_profiller/one-25/Qwen2-72B-Instruct25-input_3600-ouput_500.xlsx")
    for i in range (0,len(files)):
        print(i)
        file_name=files[i]
        file_name_a100=files_a100[i]
        print(file_name)
        print(file_name_a100)
        k_list = ['stage', 'Name', 'Instances', 'Total Time (us)', 'Avg (us)', 'Time(%)']
        c500_dir=dataset_path+file_name
        a100_dir=a100_path+file_name_a100
       
        c500_prefill, c500_decoding = filter_data(c500_dir)
        a100_prefill, a100_decoding = filter_data(a100_dir)

        c500_pre = merge_key(c500_prefill, 'mccl')
        c500_dec = merge_key(c500_decoding, 'mccl')

        a100_pre = merge_key(a100_prefill, 'nccl')
        a100_dec = merge_key(a100_decoding, 'nccl')

        prefill_ = mergea100_c500(c500_pre,a100_pre)
        decoding_ = mergea100_c500(c500_dec, a100_dec)


        dict2 = {}
        dict2['stage'] = len(prefill_[0])*['prefill']+len(decoding_[0])*['decoding']
        dict2['Name'] = prefill_[0]+ decoding_[0]
        dict2['c500_instance'] = prefill_[1]+ decoding_[1]
        dict2['c500_time'] = prefill_[2]+ decoding_[2]
        dict2['c500_per'] = prefill_[3]+ decoding_[3]

        dict2['a100_instance'] = prefill_[4]+ decoding_[4]
        dict2['a100_time'] = prefill_[5]+ decoding_[5]
        dict2['a100_per'] = prefill_[6]+ decoding_[6]
        path_result='./'+file_name
        #print(path_result)
        df = pd.DataFrame(dict2).to_excel(path_result, index=False)

    return "df"

def concurrency(dataset_path):

    all_rows = []

    files = [f for f in os.listdir(dataset_path) if f.endswith('.xlsx')]
    files.sort()
    for file_name in files:
        print(f"concurrency {file_name}")
        one_file_data = pd.read_excel(os.path.join(dataset_path, file_name))
        file_rows = process(one_file_data, file_name[:-5])
        all_rows.extend(file_rows)
    df = pd.DataFrame(all_rows)
    return df

def sort_byc500percent(df):
    # 确保列名正确（去除可能的空格）
    df.columns = df.columns.str.strip()

    # **拆分数据**
    df_prefill = df[df['stage'] == 'prefill'].copy()
    df_decoding = df[df['stage'] == 'decoding'].copy()

    # **按 `percent_c500` 降序排序**
    df_prefill_sorted = df_prefill.sort_values(by='percent_c500', ascending=False)
    df_decoding_sorted = df_decoding.sort_values(by='percent_c500', ascending=False)

    # **合并数据**
    df_sorted = pd.concat([df_prefill_sorted, df_decoding_sorted])

    # **恢复 `stage` 列格式**
    df_sorted.loc[df_sorted.duplicated(subset=['stage']), 'stage'] = ""
    return df_sorted

def change_style(file_path):
    wb = load_workbook(file_path)
    ws = wb.active

    last_column_index = ws.max_column
    last_column_name = ws.cell(row=1, column=last_column_index).value

    # 颜色填充规则
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # 绿色
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # 浅红色

    # 遍历最后一列的值，设置整行颜色 & 百分比格式
    for row in range(2, ws.max_row + 1):  # 从第二行开始，跳过表头
        cell = ws.cell(row=row, column=last_column_index)  # 获取最后一列单元格
        try:
            value = float(cell.value)  # 转换为浮点数
            cell.number_format = numbers.FORMAT_PERCENTAGE_00  # 设置百分比格式
            
            if value >= 0.7:  # 大于等于 70%
                cell.fill = green_fill  # 仅最后一列填充绿色
            else:  # 小于 70%
                for col in range(3, last_column_index + 1):  # 整行填充红色
                    ws.cell(row=row, column=col).fill = red_fill
        except (ValueError, TypeError):  # 遇到非数值数据跳过
            pass

    wb.save(file_path)

import sys
if len(sys.argv) != 2:
    print("Please set output file name!")


# file_path = f"model_cc.xlsx"
a100_path = "a100-script/result/"
c500_path = "c500-script/result/"
df2 = get_concurrency(c500_path, a100_path)
file_path = f"model_{sys.argv[1]}.xlsx"
result_path="./"
all_one=concurrency(result_path)
last_column = all_one.columns[-1]
# 过滤掉最后一列中 NaN 或 0 的行
df_filtered  = all_one[(all_one[last_column] != "NaN") & (all_one[last_column] != 0)]
df_filtered = sort_byc500percent(df_filtered)

with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
    df_filtered .to_excel(writer, sheet_name='concurrency_all', index=False)
change_style(file_path)
print(f"SUCC. File have save to {file_path}")
